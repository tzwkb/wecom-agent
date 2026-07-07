#!/usr/bin/env python3
"""读取/解析本地文件为文本 —— 文本类直读; xlsx/xls/pdf/docx 解析; 图片型内容交多模态。

返回值约定:
  - 可文本化(文本/表格/文本PDF) → 返回提取出的文本(可能截断)。
  - 需要"看"的(图片、图片型/扫描PDF、无文本docx) → 返回以 `VISUAL_MARK` 开头的标记串:
        "🖼️VISUAL <绝对路径> — <原因>（请用多模态 Read 直接看）"
    调用方(agent)应改用多模态 Read 该路径, 而不是当文本用。
  - 出错/不支持 → 返回 "[...]" 说明串。

边界已处理: 文件不存在/非文件/空文件、加密PDF、损坏文件、缺解析库、
二进制误判(NUL)、非utf8编码回退、混合PDF(文本+内嵌图)、超大(页/行/字符封顶)。

独立用: python3 read_doc.py <path>
"""
import os

VISUAL_MARK = "🖼️VISUAL"
TEXT_EXT = {"txt", "csv", "tsv", "md", "markdown", "json", "srt", "vtt", "xml",
            "log", "html", "htm", "yaml", "yml", "ini", "conf", "toml", "py", "js", "ts", "css"}
IMAGE_EXT = {"jpg", "jpeg", "png", "gif", "bmp", "webp", "heic", "heif", "tiff", "tif"}
PDF_TEXT_MIN = 50        # PDF 提取文本少于此 → 判为图片型/扫描
PAGE_CAP = 50            # PDF/最多解析页数
ROW_CAP = 400            # 表格每表最多行


def _visual(path, why):
    return f"{VISUAL_MARK} {path} — {why}（请用多模态 Read 直接看）"


def read_file(path, limit=8000):
    if not os.path.exists(path):
        return f"[文件不存在: {path}]"
    if not os.path.isfile(path):
        return f"[不是文件: {path}]"
    if os.path.getsize(path) == 0:
        return "[空文件, 0 字节]"
    base = os.path.basename(path)
    ext = base.rsplit(".", 1)[-1].lower() if "." in base else ""
    try:
        if ext in IMAGE_EXT:
            return _visual(path, f"图片({ext})")
        if ext == "pdf":
            return _pdf(path, limit)
        if ext in ("xlsx", "xlsm"):
            return _xlsx(path, limit)
        if ext == "xls":
            return _xls(path, limit)
        if ext == "docx":
            return _docx(path, limit)
        if ext in TEXT_EXT or ext == "":
            return _text(path, limit)
    except Exception as e:
        return f"[解析失败({ext or '?'}): {e}]"
    return f"[{ext or '二进制'} 格式, {os.path.getsize(path)}B — 暂不支持直读]"


def _clip(t, limit):
    return t[:limit] + (f"\n…(共 {len(t)} 字符, 已截断)" if len(t) > limit else "")


def _text(path, limit):
    with open(path, "rb") as f:
        data = f.read()
    if b"\x00" in data[:8192]:                       # 含 NUL → 八成是二进制误命名
        return f"[疑似二进制(含 NUL 字节), {len(data)}B — 不当文本读]"
    t = data.decode("utf-8", errors="replace")
    if t.count("�") > max(20, len(t) // 20):    # utf-8 替换字符过多 → 试其他编码
        for enc in ("gb18030", "big5", "latin-1"):
            try:
                cand = data.decode(enc)
            except (UnicodeDecodeError, LookupError):
                continue
            if cand.count("�") < t.count("�"):
                t = cand
                break
    return _clip(t, limit)


def _pdf(path, limit):
    try:
        import pdfplumber
    except ImportError:
        return _pdf_pypdf(path, limit)
    try:
        with pdfplumber.open(path) as pdf:
            npages = len(pdf.pages)
            pages = pdf.pages[:PAGE_CAP]
            texts = [(p.extract_text() or "") for p in pages]
            nimg = sum(len(p.images) for p in pages)
    except Exception as e:
        m = str(e).lower()
        if "password" in m or "encrypt" in m:
            return f"[加密 PDF, 无法解析: {path}]"
        return _pdf_pypdf(path, limit)               # pdfplumber 崩 → 退 pypdf 再试
    t = "\n".join(texts).strip()
    if len(t) < PDF_TEXT_MIN:                         # 文本层空 → 图片型/扫描 → 交多模态
        return _visual(path, f"图片型/扫描 PDF（{npages}页, 文本层仅 {len(t)} 字符）")
    head = f"（PDF {npages}页" + (f", 含 {nimg} 张内嵌图" if nimg else "") + "）\n"
    tail = f"\n⚠️ 此 PDF 含 {nimg} 张内嵌图, 图内信息文本提不到 → 需要时多模态 Read: {path}" if nimg else ""
    return head + _clip(t, limit) + tail


def _pdf_pypdf(path, limit):
    try:
        from pypdf import PdfReader
    except ImportError:
        return "[需 pip install pdfplumber 或 pypdf 才能解析 PDF]"
    try:
        r = PdfReader(path)
        if r.is_encrypted and r.decrypt("") == 0:
            return f"[加密 PDF, 无法解析: {path}]"
        npages = len(r.pages)
        t = "\n".join((p.extract_text() or "") for p in r.pages[:PAGE_CAP]).strip()
    except Exception as e:
        return f"[PDF 解析失败: {e}]"
    if len(t) < PDF_TEXT_MIN:
        return _visual(path, f"图片型/扫描 PDF（{npages}页）")
    return f"（PDF {npages}页）\n" + _clip(t, limit)


def _xlsx(path, limit):
    try:
        import openpyxl
    except ImportError:
        return "[需 pip install openpyxl 才能解析 xlsx]"
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        return f"[xlsx 解析失败: {e}]"
    out = []
    try:
        for ws in wb.worksheets:
            out.append(f"## 工作表「{ws.title}」")
            n = 0
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if any(cells):
                    out.append(" | ".join(cells))
                    n += 1
                if n >= ROW_CAP:
                    out.append("…(行数超上限, 截断)")
                    break
            if sum(len(x) for x in out) > limit * 3:
                out.append("…(内容过多, 截断)")
                break
    finally:
        wb.close()
    txt = "\n".join(out).strip()
    return txt or "[xlsx 无非空单元格]"


def _xls(path, limit):
    try:
        import xlrd
    except ImportError:
        return "[需 pip install xlrd 才能解析旧版 .xls]"
    try:
        wb = xlrd.open_workbook(path)
    except Exception as e:
        return f"[xls 解析失败: {e}]"
    out = []
    for sh in wb.sheets():
        out.append(f"## 工作表「{sh.name}」")
        for r in range(min(sh.nrows, ROW_CAP)):
            out.append(" | ".join(str(sh.cell_value(r, c)) for c in range(sh.ncols)))
        if sh.nrows > ROW_CAP:
            out.append("…(行数超上限, 截断)")
    txt = "\n".join(out).strip()
    return txt or "[xls 为空]"


def _docx(path, limit):
    try:
        import docx
    except ImportError:
        return "[需 pip install python-docx 才能解析 docx]"
    try:
        d = docx.Document(path)
    except Exception as e:
        return f"[docx 解析失败: {e}]"
    t = "\n".join(p.text for p in d.paragraphs if p.text).strip()
    nimg = len(d.inline_shapes)
    if len(t) < 20 and nimg:                          # 几乎无文字但有图 → 交多模态
        return _visual(path, f"docx 几乎无文本但含 {nimg} 张图")
    tail = f"\n⚠️ 含 {nimg} 张内嵌图, 需要时多模态看 → {path}" if nimg else ""
    return _clip(t, limit) + tail if t else "[docx 无文本内容]"


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        sys.exit("用法: read_doc.py <文件路径>")
    print(read_file(sys.argv[1], limit=20000))
